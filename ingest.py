"""
Read text from local files for RoPac training (ingestion).
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

CHUNK_SIZE = 1200
CHUNK_OVERLAP = 150

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".json", ".csv", ".xml", ".html", ".htm",
    ".py", ".js", ".ts", ".java", ".kt", ".kts", ".gradle", ".yaml", ".yml",
    ".sql", ".sh", ".bash", ".zsh", ".log", ".ini", ".cfg", ".toml",
    ".rst", ".tex", ".rtf",
}


def _read_plain(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ImportError("PDF support requires pypdf. Run: pip install pypdf") from e
    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text)
    return "\n\n".join(parts)


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as e:
        raise ImportError("Word support requires python-docx. Run: pip install python-docx") from e
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _read_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ImportError("Excel support requires openpyxl. Run: pip install openpyxl") from e
    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"## Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(c.strip() for c in cells):
                parts.append("\t".join(cells))
    wb.close()
    return "\n".join(parts)


def _read_csv(path: Path) -> str:
    rows = []
    with path.open(newline="", encoding="utf-8", errors="replace") as f:
        reader = csv.reader(f)
        for row in reader:
            rows.append("\t".join(row))
    return "\n".join(rows)


def read_file(path: Path) -> str:
    path = path.expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")
    if not path.is_file():
        raise ValueError(f"Not a file: {path}")

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix in {".docx"}:
        return _read_docx(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx(path)
    if suffix == ".csv":
        return _read_csv(path)
    if suffix in TEXT_EXTENSIONS or suffix == "":
        return _read_plain(path)

    raise ValueError(
        f"Unsupported file type: {suffix or '(no extension)'}. "
        "Supported: pdf, docx, xlsx, csv, txt, md, code files, json, etc."
    )


def chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    text = re_sub_whitespace(text)
    if not text:
        return []
    if len(text) <= size:
        return [text]

    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = start + size
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start = end - overlap
    return chunks


def re_sub_whitespace(text: str) -> str:
    import re
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
